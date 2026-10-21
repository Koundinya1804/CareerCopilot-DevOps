"""
app/utils/pdf_parser.py
───────────────────────
PDF and DOCX text extraction with section detection.

Supports:
  - pypdf   : pure-Python PDF parser (no system deps)
  - python-docx : DOCX extraction
  - Section detection via heading heuristics
  - Metadata extraction (pages, author, title)
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.core.logging import get_logger

logger = get_logger(__name__)

# Known resume section headings — used for structural segmentation
_SECTION_PATTERNS: list[tuple[str, str]] = [
    (r"(?i)^(summary|professional\s+summary|objective|profile|about\s+me)", "summary"),
    (r"(?i)^(experience|work\s+experience|employment|work\s+history|professional\s+experience)", "experience"),
    (r"(?i)^(education|academic|qualifications)", "education"),
    (r"(?i)^(skills|technical\s+skills|core\s+competencies|technologies|tools)", "skills"),
    (r"(?i)^(projects|personal\s+projects|academic\s+projects|portfolio)", "projects"),
    (r"(?i)^(certifications?|certificates?|licen[sc]es?|credentials)", "certifications"),
    (r"(?i)^(awards?|achievements?|honors?|accomplishments?)", "achievements"),
    (r"(?i)^(publications?|research|papers?)", "publications"),
    (r"(?i)^(languages?)", "languages"),
    (r"(?i)^(interests?|hobbies|activities)", "interests"),
    (r"(?i)^(volunteer|community|social)", "volunteer"),
    (r"(?i)^(contact|personal\s+info)", "contact"),
]


@dataclass
class ParsedResume:
    """Structured result from text extraction."""
    raw_text:    str = ""
    page_count:  int = 0
    word_count:  int = 0
    char_count:  int = 0
    sections:    dict[str, str] = field(default_factory=dict)
    metadata:    dict[str, str] = field(default_factory=dict)
    mime_type:   str = "application/pdf"

    @property
    def is_empty(self) -> bool:
        return len(self.raw_text.strip()) < 50

    @property
    def has_section(self) -> bool:
        return bool(self.sections)


# ── PDF Extraction ─────────────────────────────────────────────────────────────
def extract_from_pdf(file_bytes: bytes) -> ParsedResume:
    """Extract text from PDF bytes using pypdf."""
    try:
        import pypdf  # type: ignore[import]
    except ImportError:
        raise RuntimeError("pypdf not installed. Run: pip install pypdf")

    result = ParsedResume(mime_type="application/pdf")

    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        result.page_count = len(reader.pages)

        # Extract metadata
        if reader.metadata:
            meta = reader.metadata
            result.metadata = {
                "title":    str(meta.title  or ""),
                "author":   str(meta.author or ""),
                "creator":  str(meta.creator or ""),
            }

        # Extract text from every page
        pages_text: list[str] = []
        for page_num, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text(extraction_mode="layout") or ""
                if not text.strip():
                    # Fallback to plain extraction
                    text = page.extract_text() or ""
                pages_text.append(text)
                logger.debug("pdf_page_extracted", page=page_num, chars=len(text))
            except Exception as e:
                logger.warning("pdf_page_extraction_failed", page=page_num, error=str(e))
                pages_text.append("")

        result.raw_text = "\n\n".join(pages_text)
        result.word_count  = len(result.raw_text.split())
        result.char_count  = len(result.raw_text)
        result.sections    = _detect_sections(result.raw_text)

        logger.info(
            "pdf_extracted",
            pages=result.page_count,
            words=result.word_count,
            sections=list(result.sections.keys()),
        )
        return result

    except Exception as exc:
        logger.error("pdf_extraction_error", error=str(exc))
        raise ValueError(f"Could not extract text from PDF: {exc}") from exc


# ── DOCX Extraction ────────────────────────────────────────────────────────────
def extract_from_docx(file_bytes: bytes) -> ParsedResume:
    """Extract text from DOCX bytes using python-docx."""
    try:
        import docx  # type: ignore[import]
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    result = ParsedResume(mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract from tables (skills tables etc.)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        result.raw_text   = "\n".join(paragraphs)
        result.page_count = max(1, len(result.raw_text) // 3000)  # estimate
        result.word_count = len(result.raw_text.split())
        result.char_count = len(result.raw_text)
        result.sections   = _detect_sections(result.raw_text)

        logger.info("docx_extracted", words=result.word_count, sections=list(result.sections.keys()))
        return result

    except Exception as exc:
        logger.error("docx_extraction_error", error=str(exc))
        raise ValueError(f"Could not extract text from DOCX: {exc}") from exc


# ── Dispatcher ────────────────────────────────────────────────────────────────
def extract_text(file_bytes: bytes, mime_type: str) -> ParsedResume:
    """Route to the correct extractor based on MIME type."""
    mime = mime_type.lower()
    if "pdf" in mime:
        return extract_from_pdf(file_bytes)
    elif "wordprocessingml" in mime or "docx" in mime or "msword" in mime:
        return extract_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {mime_type}")


# ── Section Detection ─────────────────────────────────────────────────────────
def _detect_sections(text: str) -> dict[str, str]:
    """
    Split resume text into named sections using heading heuristics.
    Returns dict mapping section_name → section_content.
    """
    lines = text.split("\n")
    sections: dict[str, str] = {}
    current_section = "header"
    current_lines: list[str] = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            current_lines.append("")
            continue

        matched_section = _match_section_heading(stripped)
        if matched_section and len(stripped) < 60:
            # Save previous section
            if current_lines:
                sections[current_section] = "\n".join(current_lines).strip()
            current_section = matched_section
            current_lines = []
        else:
            current_lines.append(line)

    # Save final section
    if current_lines:
        sections[current_section] = "\n".join(current_lines).strip()

    return sections


def _match_section_heading(line: str) -> str | None:
    """Return section key if line matches a known heading, else None."""
    for pattern, key in _SECTION_PATTERNS:
        if re.match(pattern, line):
            return key
    return None


# ── Text Cleaning ─────────────────────────────────────────────────────────────
def clean_text(text: str) -> str:
    """Normalise whitespace and remove non-printable characters."""
    text = re.sub(r"[^\x20-\x7E\n]", " ", text)   # keep printable ASCII + newline
    text = re.sub(r" {2,}", " ", text)              # collapse multiple spaces
    text = re.sub(r"\n{3,}", "\n\n", text)          # max 2 consecutive newlines
    return text.strip()
