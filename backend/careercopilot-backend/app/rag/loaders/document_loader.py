"""
app/rag/loaders/document_loader.py
────────────────────────────────────
Loads documents from files or directories into LangChain Document objects.

Supported formats:
  .pdf   → pypdf-based extraction
  .docx  → python-docx
  .txt   → plain UTF-8
  .md    → Markdown (treated as plain text)
  .json  → structured Q&A pairs ({"question": ..., "answer": ...})
  .jsonl → one JSON object per line

Each loaded Document carries rich metadata:
  source, category, filename, file_type, page (if applicable),
  char_count, word_count, loaded_at
"""
from __future__ import annotations

import json
from datetime import datetime, UTC
from pathlib import Path
from typing import Iterator

from langchain_core.documents import Document

from app.core.logging import get_logger

logger = get_logger(__name__)

# ── Type alias ────────────────────────────────────────────────────────────────
LoadedDoc = Document


# ── Base metadata builder ─────────────────────────────────────────────────────
def _base_meta(path: Path, category: str, extra: dict | None = None) -> dict:
    meta = {
        "source":    str(path),
        "filename":  path.name,
        "file_type": path.suffix.lower().lstrip("."),
        "category":  category,
        "loaded_at": datetime.now(UTC).isoformat(),
    }
    if extra:
        meta.update(extra)
    return meta


def _enrich(doc: Document) -> Document:
    """Add word/char counts to metadata after content is set."""
    doc.metadata["word_count"] = len(doc.page_content.split())
    doc.metadata["char_count"] = len(doc.page_content)
    return doc


# ── Format-specific loaders ───────────────────────────────────────────────────

def _load_pdf(path: Path, category: str) -> list[LoadedDoc]:
    """Extract per-page Documents from a PDF."""
    try:
        import pypdf  # type: ignore[import]
    except ImportError:
        raise RuntimeError("pypdf not installed: pip install pypdf")

    docs: list[LoadedDoc] = []
    reader = pypdf.PdfReader(str(path))

    for page_num, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
            text = text.strip()
        except Exception:
            text = ""

        if not text:
            continue

        doc = Document(
            page_content=text,
            metadata=_base_meta(path, category, {
                "page":       page_num,
                "total_pages": len(reader.pages),
            }),
        )
        docs.append(_enrich(doc))

    logger.info("pdf_loaded", file=path.name, pages=len(docs), category=category)
    return docs


def _load_docx(path: Path, category: str) -> list[LoadedDoc]:
    try:
        import docx  # type: ignore[import]
    except ImportError:
        raise RuntimeError("python-docx not installed: pip install python-docx")

    doc_obj = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in doc_obj.paragraphs if p.text.strip()]

    # Also grab table content
    for table in doc_obj.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n".join(paragraphs)
    if not full_text.strip():
        return []

    doc = Document(
        page_content=full_text,
        metadata=_base_meta(path, category),
    )
    return [_enrich(doc)]


def _load_text(path: Path, category: str) -> list[LoadedDoc]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    doc = Document(page_content=text, metadata=_base_meta(path, category))
    return [_enrich(doc)]


def _load_json(path: Path, category: str) -> list[LoadedDoc]:
    """
    Load structured JSON knowledge files.

    Supported formats:
      List of {"question": ..., "answer": ...}
      List of {"topic": ..., "content": ...}
      Single dict with "content" key
      Any list of dicts → joined as key: value pairs
    """
    raw = json.loads(path.read_text(encoding="utf-8"))
    items: list[dict] = []

    if isinstance(raw, list):
        items = raw
    elif isinstance(raw, dict):
        items = [raw]
    else:
        return []

    docs: list[LoadedDoc] = []
    for i, item in enumerate(items):
        if not isinstance(item, dict):
            continue

        # Q&A format
        if "question" in item and "answer" in item:
            text = f"Q: {item['question']}\nA: {item['answer']}"
            extra = {"item_index": i, "format": "qa"}
        elif "topic" in item and "content" in item:
            text = f"Topic: {item['topic']}\n\n{item['content']}"
            extra = {"item_index": i, "format": "topic", "topic": item["topic"]}
        elif "content" in item:
            text = str(item["content"])
            extra = {"item_index": i, "format": "content"}
        else:
            # Generic: join all string values
            parts = [f"{k}: {v}" for k, v in item.items() if isinstance(v, str)]
            text = "\n".join(parts)
            extra = {"item_index": i, "format": "generic"}

        if text.strip():
            doc = Document(page_content=text, metadata=_base_meta(path, category, extra))
            docs.append(_enrich(doc))

    logger.info("json_loaded", file=path.name, items=len(docs), category=category)
    return docs


def _load_jsonl(path: Path, category: str) -> list[LoadedDoc]:
    docs: list[LoadedDoc] = []
    for i, line in enumerate(path.read_text(encoding="utf-8").splitlines()):
        line = line.strip()
        if not line:
            continue
        try:
            item = json.loads(line)
            sub_docs = _load_json.__wrapped__(path, category) if hasattr(_load_json, "__wrapped__") else []
            # Inline: process single item
            if "question" in item and "answer" in item:
                text = f"Q: {item['question']}\nA: {item['answer']}"
            elif "content" in item:
                text = str(item["content"])
            else:
                text = " ".join(str(v) for v in item.values() if isinstance(v, str))
            if text.strip():
                doc = Document(
                    page_content=text,
                    metadata=_base_meta(path, category, {"item_index": i, "format": "jsonl"}),
                )
                docs.append(_enrich(doc))
        except json.JSONDecodeError:
            continue
    return docs


# ── Dispatcher ────────────────────────────────────────────────────────────────
_LOADERS = {
    ".pdf":   _load_pdf,
    ".docx":  _load_docx,
    ".doc":   _load_docx,
    ".txt":   _load_text,
    ".md":    _load_text,
    ".json":  _load_json,
    ".jsonl": _load_jsonl,
}


def load_file(path: str | Path, category: str) -> list[LoadedDoc]:
    """Load a single file into Documents. Returns [] for unsupported types."""
    p = Path(path)
    if not p.exists():
        logger.warning("file_not_found", path=str(p))
        return []

    loader_fn = _LOADERS.get(p.suffix.lower())
    if not loader_fn:
        logger.warning("unsupported_file_type", suffix=p.suffix, file=p.name)
        return []

    try:
        return loader_fn(p, category)
    except Exception as exc:
        logger.error("file_load_error", file=str(p), error=str(exc))
        return []


def load_directory(
    directory: str | Path,
    category:  str,
    recursive: bool = True,
) -> Iterator[LoadedDoc]:
    """
    Walk a directory and yield Documents from all supported files.
    Skips hidden files and __pycache__ directories.
    """
    root = Path(directory)
    if not root.exists():
        logger.warning("directory_not_found", path=str(root))
        return

    pattern = "**/*" if recursive else "*"
    for fp in sorted(root.glob(pattern)):
        if fp.is_dir():
            continue
        if fp.name.startswith(".") or "__pycache__" in fp.parts:
            continue
        docs = load_file(fp, category)
        for doc in docs:
            yield doc


def load_bytes(
    data:      bytes,
    filename:  str,
    category:  str,
    mime_type: str = "application/pdf",
) -> list[LoadedDoc]:
    """Load documents from in-memory bytes (e.g. user-uploaded files)."""
    import tempfile
    suffix = Path(filename).suffix or ".pdf"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(data)
        tmp_path = Path(tmp.name)
    try:
        return load_file(tmp_path, category)
    finally:
        tmp_path.unlink(missing_ok=True)
